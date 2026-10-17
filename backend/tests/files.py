"""Tiny but structurally valid documents for upload tests."""

import io
import zipfile

PDF_BYTES = (
    b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    b"2 0 obj<</Type/Pages/Count 0/Kids[]>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n"
)


def docx_bytes(text: str = "Emergency repairs are attended within 24 hours.") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
            'content-types"><Default Extension="xml" ContentType="application/xml"/></Types>',
        )
        z.writestr(
            "word/document.xml",
            '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
            f'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p>'
            "</w:body></w:document>",
        )
    return buf.getvalue()


def plain_zip_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("readme.txt", "not a word document")
    return buf.getvalue()


def pdf_with_text(pages: list[list[str]]) -> bytes:
    """A real text PDF: each page is a list of lines; "" makes a blank line between paragraphs."""

    def esc(line: str) -> str:
        return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: list[bytes] = []
    page_count = len(pages)
    font_id = 3 + 2 * page_count
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(page_count))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode())
    for i, lines in enumerate(pages):
        content_id = 4 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 740 Td"]
        for line in lines:
            ops.append(f"({esc(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def docx_with_sections(sections: dict[str, list[str]]) -> bytes:
    """A real Word document with Heading 1 sections and body paragraphs."""
    import docx

    document = docx.Document()
    for heading, paragraphs in sections.items():
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
